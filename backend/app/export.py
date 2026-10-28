"""Read a .docx resume as text, and write a Resume back out as a clean .docx."""

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

from .models import Resume

NAVY = RGBColor(0x1F, 0x38, 0x64)
FONT = "Calibri"


def docx_to_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:  # some templates put the header in a table
        for row in table.rows:
            lines.append(" | ".join(c.text for c in row.cells if c.text.strip()))
    return "\n".join(lines)


def _bottom_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, val in {"val": "single", "sz": "6", "space": "1", "color": "1F3864"}.items():
        bottom.set(qn(f"w:{key}"), val)
    borders.append(bottom)
    p_pr.append(borders)


def _para(doc, text="", *, bold=False, size=10, align=None, space_after=2, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if align:
        p.alignment = align
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
    return p


def _heading(doc, text: str) -> None:
    p = _para(doc, text.upper(), bold=True, size=12, space_after=4, color=NAVY)
    p.paragraph_format.space_before = Pt(8)
    _bottom_border(p)


def _bullet(doc, label: str, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    if label:
        r = p.add_run(f"{label}: ")
        r.bold = True
        r.font.size = Pt(10)
    r = p.add_run(text)
    r.font.size = Pt(10)


def resume_to_docx(resume: Resume) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Twips(12240), Twips(15840)  # US Letter
    section.top_margin = section.bottom_margin = Twips(1008)
    section.left_margin = section.right_margin = Twips(1080)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)

    _para(doc, resume.name.upper(), bold=True, size=24, align=WD_ALIGN_PARAGRAPH.CENTER)
    if resume.contact:
        _para(doc, " | ".join(resume.contact), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

    if resume.profile:
        _heading(doc, "Profile")
        p = _para(doc, resume.profile)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if resume.skills:
        _heading(doc, "Technical Skills")
        for group in resume.skills:
            p = _para(doc)
            r = p.add_run(f"{group.name}: ")
            r.bold = True
            p.add_run(" | ".join(group.items))

    if resume.experience:
        _heading(doc, "Professional Experience")
        for role in resume.experience:
            p = _para(doc, space_after=2)
            p.paragraph_format.space_before = Pt(4)
            p.add_run(role.title).bold = True
            dates = " – ".join(x for x in (role.start, role.end) if x)
            meta = " | ".join(x for x in (role.company, role.location, dates) if x)
            p.add_run(f" | {meta}").italic = True
            for b in role.bullets:
                _bullet(doc, b.label, b.text)

    for title, items in (("Certifications", resume.certifications), ("Education", resume.education)):
        if items:
            _heading(doc, title)
            for item in items:
                _bullet(doc, "", item)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
